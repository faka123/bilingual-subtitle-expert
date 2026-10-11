import unittest
from io import BytesIO
from docx import Document
from subtitle_app.core import (
    parse_srt, parse_proofread_doc, parse_docx, decode_text_file,
    process_subtitles, process_bilingual_api, quality_check, split_english_text,
    split_english_text_llm, ProofreadPair, TimelineItem,
)


def cue(index, start, text, end='00:00:05,000'):
    return f'{index}\n{start} --> {end}\n{text}'


class SubtitleTests(unittest.TestCase):
    def test_file_encodings_and_bom_keep_first_cue(self):
        text = cue(1, '00:00:01,000', '你好世界')
        for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
            self.assertEqual(parse_srt(decode_text_file(text.encode(encoding)))[0].text, '你好世界')
        self.assertEqual(len(parse_srt('\ufeff' + text)), 1)

    def test_incomplete_input_is_not_successful(self):
        with self.assertRaisesRegex(ValueError, '中英对照'):
            process_subtitles('只有中文', cue(1, '00:00:01,000', '你好'))
        with self.assertRaisesRegex(ValueError, '时间轴格式'):
            parse_srt('1\ninvalid\n你好')

    def test_preserve_empty_abnormal_duplicate_and_original_numbers(self):
        timeline = '\n\n'.join([
            cue(8, '00:00:04.000', '你好世界'),
            cue(8, '00:00:02,000', '再见世界'),
            cue(96, '00:00:70,000', '', '00:00:01,000'),
        ])
        output, stats = process_subtitles('你好世界\nHello world\n\n再见世界\nGoodbye world', timeline)
        entries = parse_srt(output)
        originals = parse_srt(timeline)
        self.assertEqual(entries[:3], originals)
        self.assertEqual([e.index for e in entries[3:]], [97, 98, 99])
        self.assertEqual([e.start for e in entries[3:]], [e.start for e in originals])
        self.assertEqual([e.text for e in entries[3:]], ['Hello world', 'Goodbye world', ''])
        self.assertEqual(stats['unmatched_entries'], 1)
        self.assertTrue(any('时间轴数值异常' in i for i in quality_check(output, stats)))

    def test_interleaved_paragraphs_keep_input_order(self):
        timeline = '\n\n'.join(cue(i, f'00:00:0{i},000', text)
                               for i, text in enumerate(['苹果上市', '香蕉成熟', '苹果成熟'], 1))
        output, _ = process_subtitles('苹果上市苹果成熟\nApple launch Apple ripe\n\n香蕉成熟\nBanana ripe', timeline)
        entries = parse_srt(output)
        self.assertEqual([e.start for e in entries[:3]], [e.start for e in entries[3:]])
        self.assertEqual(entries[4].text, 'Banana ripe')
        self.assertEqual(' '.join([entries[3].text, entries[5].text]).split(), 'Apple launch Apple ripe'.split())

    def test_old_english_is_ignored(self):
        timeline = cue(1, '00:00:01,000', '你好世界') + '\n\n' + cue(2, '00:00:01,000', 'Wrong old translation')
        output, stats = process_subtitles('你好世界\nHello world', timeline)
        self.assertNotIn('Wrong', output)
        self.assertEqual(len(parse_srt(output)), 2)
        self.assertEqual(stats['srt_en_entries_filtered'], 1)

    def test_unrelated_cue_is_not_forced_to_first_paragraph(self):
        output, stats = process_subtitles('你好世界\nHello world', cue(1, '00:00:01,000', '火星能源革命'))
        self.assertEqual(parse_srt(output)[1].text, '')
        self.assertEqual(stats['unmatched_entries'], 1)
        self.assertEqual(stats['unused_paragraphs'], 1)

    def test_short_english_does_not_duplicate_words(self):
        parts = split_english_text('Hello world', 4, ['你', '好', '世', '界'])
        self.assertEqual(len(parts), 4)
        self.assertEqual(' '.join(parts).split(), ['Hello', 'world'])

    def test_ai_missing_words_or_reordering_falls_back(self):
        class FakeLLM:
            def split_english_text(self, text, count):
                return ['world', 'Hello']
        errors = []
        parts = split_english_text_llm('Hello world', 2, FakeLLM(), errors=errors)
        self.assertEqual(' '.join(parts).split(), ['Hello', 'world'])
        self.assertEqual(len(errors), 1)

    def test_ai_match_does_not_enable_ai_split(self):
        class FakeLLM:
            def match_subtitles(self, texts, paragraphs):
                return [{'srt_index': i, 'paragraph_index': 0, 'confidence': 1} for i in range(len(texts))]
            def split_english_text(self, text, count):
                raise AssertionError('AI split must not be called')
        output, stats = process_subtitles('你好世界\nHello world', cue(1, '00:00:01,000', '你好世界'),
                                         llm_service=FakeLLM(), use_llm_match=True, use_llm_split=False)
        self.assertIn('Hello world', output)
        self.assertEqual(stats['llm_split_count'], 0)

    def test_punctuation_only_changes_english_end(self):
        output, _ = process_subtitles('你好世界\nMeasure 0.025/mg³, please!', cue(1, '00:00:01,000', '你好世界。'),
                                       remove_end_punctuation=True)
        entries = parse_srt(output)
        self.assertEqual(entries[0].text, '你好世界。')
        self.assertEqual(entries[1].text, 'Measure 0.025/mg³, please')

    def test_markup_merged_line_and_docx_body_order(self):
        pairs = parse_proofread_doc("你好世界Hello \\'world\\' {.mark}\n\n再见世界\n[~~Goodbye world~~]{.mark}")
        self.assertEqual([(p.chinese, p.english) for p in pairs], [('你好世界', "Hello 'world'"), ('再见世界', 'Goodbye world')])
        doc = Document()
        doc.add_paragraph('你好世界')
        doc.add_paragraph('Hello world')
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = '再见世界'
        row.cells[1].text = 'Goodbye world'
        doc.add_paragraph('苹果成熟')
        doc.add_paragraph('Apple ripe')
        buffer = BytesIO()
        doc.save(buffer)
        pairs = parse_proofread_doc(parse_docx(buffer.getvalue()))
        self.assertEqual([p.english for p in pairs], ['Hello world', 'Goodbye world', 'Apple ripe'])

    def test_api_order_and_original_chinese(self):
        items = [TimelineItem(9, '00:00:02,000', '00:00:01,000', '你好世界'),
                 TimelineItem(9, '00:00:03,000', '00:00:04,000', '再见世界')]
        result = process_bilingual_api([ProofreadPair('你好世界', 'Hello world'), ProofreadPair('再见世界', 'Goodbye world')], items)
        self.assertEqual(result.chinese_part, items)
        self.assertEqual([e.index for e in result.english_part], [10, 11])
        self.assertEqual([e.text for e in result.english_part], ['Hello world', 'Goodbye world'])


if __name__ == '__main__':
    unittest.main()
